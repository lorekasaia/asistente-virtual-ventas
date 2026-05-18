import os
import uuid
import re
import logging
import pandas as pd
import matplotlib.pyplot as plt
from google import adk
import sqlalchemy
from database import obtener_motor_bd, consultar_cloud_sql, MAPA_ESTADOS

logger = logging.getLogger("BatiaAgent")

def obtener_resumen_pipeline() -> str:
    try:
        # Delegamos la agregación a la base de datos (Pushdown) en lugar de descargar toda la tabla a RAM
        engine, _ = obtener_motor_bd()
        query = sqlalchemy.text("""
            SELECT _estado, 
                   COUNT(id) as cantidad_clientes, 
                   SUM(valor_estimado) as valor_total_mxn, 
                   AVG(valor_estimado) as ticket_promedio 
            FROM clientes 
            GROUP BY _estado
        """)
        with engine.connect() as conn:
            rows = conn.execute(query).fetchall()
            
        if not rows:
            return "No hay datos suficientes para generar un resumen."
            
        datos = []
        for row in rows:
            estado_id = row[0]
            estado_texto = MAPA_ESTADOS.get(estado_id, 'Desconocido') if pd.notna(estado_id) else 'Desconocido'
            datos.append({"estado_texto": estado_texto, "cantidad_clientes": row[1], "valor_total_mxn": float(row[2]) if row[2] else 0.0, "ticket_promedio": float(row[3]) if row[3] else 0.0})
            
        df_resumen = pd.DataFrame(datos)
        return "Resumen Financiero del Pipeline:\n" + df_resumen.to_json(orient="records", force_ascii=False)
    except Exception as e:
        logger.error(f"Error en obtener_resumen_pipeline: {e}", exc_info=True)
        return f"Error al generar el resumen estadístico: {e}"

def consultar_dashboard_bi(kpi: str, contexto: str = "general") -> str:
    print(f"[BI API Mock] Solicitando KPI: '{kpi}' con contexto '{contexto}'")
    if kpi == "ventas_totales":
        return f"El Dashboard de BI reporta que las ventas totales para '{contexto}' son de $1,450,000 MXN este trimestre."
    elif kpi == "tasa_conversion":
        return f"Según la plataforma de BI, la tasa de conversión de prospectos a clientes para '{contexto}' se sitúa en un 24.5%."
    elif kpi == "rendimiento_vendedores":
        return "El reporte de BI indica que Lore lidera las ventas con un 120% de alcance de cuota, seguida de cerca por el resto del equipo."
    else:
        return f"Datos del Dashboard para el KPI '{kpi}': Los indicadores están estables y dentro de los rangos esperados para el período actual."

def generar_grafico_analisis(metrica: str) -> str:
    """
    Genera un gráfico visual basado en los datos de clientes.
    El parámetro 'metrica' debe ser una de las siguientes opciones (o un sinónimo cercano):
    - 'estado': Gráfico circular de la distribución de clientes por estado en el pipeline.
    - 'prioridad': Gráfico de barras de clientes por nivel de prioridad.
    - 'valor': Gráfico de barras del valor estimado del pipeline por estado.
    - 'fuente': Gráfico de barras horizontales del origen de los prospectos.
    - 'conversion': Gráfico circular de la tasa de conversión general.
    """
    try:
        # --- Normalización de métrica para flexibilidad ---
        # Permite que el usuario pida "estado actual" y se mapee a "estado".
        metrica_limpia = metrica.lower()
        if "estado" in metrica_limpia:
            metrica_limpia = "estado"
        elif "prioridad" in metrica_limpia:
            metrica_limpia = "prioridad"
        elif "valor" in metrica_limpia:
            metrica_limpia = "valor"
        elif "fuente" in metrica_limpia or "origen" in metrica_limpia:
            metrica_limpia = "fuente"
        elif "conversión" in metrica_limpia or "conversion" in metrica_limpia:
            metrica_limpia = "conversion"
        # --- Fin de la normalización ---

        # Usamos la API orientada a objetos (fig, ax) para evitar colisiones entre usuarios concurrentes
        fig, ax = plt.subplots(figsize=(8, 6))
        engine, _ = obtener_motor_bd()
        
        with engine.connect() as conn:
            if metrica_limpia == "prioridad":
                query = sqlalchemy.text("SELECT prioridad, COUNT(id) as conteo FROM clientes WHERE prioridad IS NOT NULL GROUP BY prioridad")
                df = pd.read_sql(query, conn)
                if df.empty: return "No hay datos de prioridad para graficar."
                conteo = df.set_index('prioridad')['conteo']
                conteo.plot(kind='bar', color=['#4CAF50', '#FF9800', '#F44336'], ax=ax)
                ax.set_title("Distribución de Clientes por Prioridad")
                ax.set_xlabel("Nivel de Prioridad")
                ax.set_ylabel("Cantidad de Clientes")
                ax.tick_params(axis='x', rotation=0)
            elif metrica_limpia == "estado":
                query = sqlalchemy.text("SELECT _estado, COUNT(id) as conteo FROM clientes WHERE _estado IS NOT NULL GROUP BY _estado")
                df = pd.read_sql(query, conn)
                if df.empty: return "No hay datos de estado para graficar."
                df['estado_texto'] = df['_estado'].map(MAPA_ESTADOS).fillna('Otro')
                conteo = df.set_index('estado_texto')['conteo']
                conteo.plot(kind='pie', autopct='%1.1f%%', startangle=90, ax=ax)
                ax.set_title("Proporción de Clientes por Estado Interno")
                ax.set_ylabel("")
            elif metrica_limpia == "valor":
                query = sqlalchemy.text("SELECT _estado, SUM(valor_estimado) as suma_valor FROM clientes WHERE _estado IS NOT NULL AND valor_estimado IS NOT NULL GROUP BY _estado")
                df = pd.read_sql(query, conn)
                if df.empty: return "No hay datos de valor estimado para graficar."
                df['estado_texto'] = df['_estado'].map(MAPA_ESTADOS).fillna('Otro')
                suma_valor = df.set_index('estado_texto')['suma_valor'].sort_values(ascending=False)
                suma_valor.plot(kind='bar', color='#2196F3', ax=ax)
                ax.set_title("Valor Estimado ($) del Pipeline por Estado")
                ax.set_xlabel("Estado del Cliente")
                ax.set_ylabel("Valor Estimado Total")
                ax.tick_params(axis='x', rotation=45)
            elif metrica_limpia == "fuente":
                query = sqlalchemy.text("SELECT COALESCE(fuente, 'Desconocido') as fuente_texto, COUNT(id) as conteo FROM clientes GROUP BY fuente")
                df = pd.read_sql(query, conn)
                if df.empty: return "No hay datos de fuentes para graficar."
                conteo = df.set_index('fuente_texto')['conteo']
                conteo.sort_values().plot(kind='barh', color='#9C27B0', ax=ax)
                ax.set_title("Origen de los Prospectos (Fuentes)")
                ax.set_xlabel("Cantidad de Clientes")
                ax.set_ylabel("Fuente")
            elif metrica_limpia == "conversion":
                query = sqlalchemy.text("SELECT es_cliente, COUNT(id) as conteo FROM clientes WHERE es_cliente IS NOT NULL GROUP BY es_cliente")
                df = pd.read_sql(query, conn)
                if df.empty: return "No hay datos de conversión para graficar."
                df['tipo'] = df['es_cliente'].map({True: 'Cliente Cerrado', False: 'Prospecto Activo'})
                conteo = df.set_index('tipo')['conteo']
                conteo.plot(kind='pie', autopct='%1.1f%%', startangle=90, colors=['#00BCD4', '#FFC107'], ax=ax)
                ax.set_title("Tasa de Conversión General")
                ax.set_ylabel("")
            else:
                return f"No se pudo generar el gráfico para la métrica '{metrica}'. Las métricas permitidas son: 'prioridad', 'estado', 'valor', 'fuente' y 'conversion'. Si el usuario pide 'todas', debes ejecutar esta herramienta 5 veces seguidas (una por cada métrica permitida)."
        
        filename = f"grafico_{uuid.uuid4().hex[:8]}.png"
        filepath = os.path.join("graficos", filename)
        fig.savefig(filepath, bbox_inches='tight')
        plt.close(fig) # Se asegura de liberar los recursos solo de esta figura
        return f"Gráfico generado con éxito. DEBES responder esto al usuario exactamente así para que vea la imagen: <br><img src='/graficos/{filename}' alt='Gráfico de {metrica_limpia}' style='max-width: 100%; border-radius: 8px; margin-top: 10px;'/>"
    except Exception as e:
        logger.error(f"Error en generar_grafico_analisis: {e}", exc_info=True)
        return f"Error al procesar y graficar los datos: {e}"

def exportar_datos_excel(termino_busqueda: str = "") -> str:
    try:
        # Limitamos la exportación a 5,000 registros para prevenir colapso de RAM y generar Excels eficientes
        df = consultar_cloud_sql(termino_busqueda, limite=5000)
        if df.empty:
            return "No hay datos en la base de datos que coincidan con ese criterio para exportar."
        
        filename = f"reporte_clientes_{uuid.uuid4().hex[:8]}.xlsx"
        filepath = os.path.join("reportes", filename)
        df.to_excel(filepath, index=False)
        return f"Reporte Excel generado con éxito. DEBES responder esto al usuario para que pueda descargarlo: <br><a href='/reportes/{filename}' download style='display: inline-block; padding: 10px 15px; background: #107c41; color: white; text-decoration: none; border-radius: 5px; margin-top: 10px; font-weight: bold;'>📊 Descargar Reporte Excel</a>"
    except Exception as e:
        logger.error(f"Error en exportar_datos_excel: {e}", exc_info=True)
        return f"Error al generar el archivo Excel: {e}"

def generar_reporte_pdf(titulo: str, contenido: str, nombre_imagen: str = "") -> str:
    """
    Crea un documento PDF con un título, un contenido de texto y opcionalmente una imagen.
    Si generaste un gráfico antes, pasa su nombre de archivo (ej. 'grafico_123.png') en 'nombre_imagen'.
    Guarda el archivo en la carpeta 'reportes' y devuelve un enlace de descarga.
    """
    try:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", 'B', size=16)
        pdf.cell(0, 10, txt=titulo, ln=True, align='C')
        pdf.ln(10)

        if nombre_imagen:
            match = re.search(r'grafico_[a-zA-Z0-9]+\.png', nombre_imagen)
            if match:
                ruta_img = os.path.join("graficos", match.group(0))
                if os.path.exists(ruta_img):
                    pdf.image(ruta_img, w=160)
                    pdf.ln(10)

        pdf.set_font("Arial", size=12)
        # El texto debe estar codificado para evitar errores con caracteres especiales en FPDF
        contenido_encoded = contenido.encode('latin-1', 'replace').decode('latin-1')
        pdf.multi_cell(0, 10, txt=contenido_encoded)
        
        filename = f"reporte_{uuid.uuid4().hex[:8]}.pdf"
        filepath = os.path.join("reportes", filename)
        pdf.output(filepath)
        
        return f"Reporte PDF generado. Responde con este enlace para descarga: <br><a href='/reportes/{filename}' download style='display: inline-block; padding: 10px 15px; background: #D32F2F; color: white; text-decoration: none; border-radius: 5px; margin-top: 10px; font-weight: bold;'>📄 Descargar Reporte PDF</a>"
    except Exception as e:
        logger.error(f"Error en generar_reporte_pdf: {e}", exc_info=True)
        return f"Error al generar el archivo PDF: {e}"

def generar_reporte_word(titulo: str, contenido: str, nombre_imagen: str = "") -> str:
    """
    Crea un documento Word (.docx) con un título, un contenido de texto y opcionalmente una imagen.
    Si generaste un gráfico antes, pasa su nombre de archivo (ej. 'grafico_123.png') en 'nombre_imagen'.
    Guarda el archivo en la carpeta 'reportes' y devuelve un enlace de descarga.
    """
    try:
        from docx import Document
        from docx.shared import Inches
        document = Document()
        document.add_heading(titulo, level=1)

        if nombre_imagen:
            match = re.search(r'grafico_[a-zA-Z0-9]+\.png', nombre_imagen)
            if match:
                ruta_img = os.path.join("graficos", match.group(0))
                if os.path.exists(ruta_img):
                    document.add_picture(ruta_img, width=Inches(6.0))

        document.add_paragraph(contenido)
        
        filename = f"reporte_{uuid.uuid4().hex[:8]}.docx"
        filepath = os.path.join("reportes", filename)
        document.save(filepath)
        
        return f"Reporte Word generado. Responde con este enlace para descarga: <br><a href='/reportes/{filename}' download style='display: inline-block; padding: 10px 15px; background: #2B579A; color: white; text-decoration: none; border-radius: 5px; margin-top: 10px; font-weight: bold;'>📄 Descargar Reporte Word</a>"
    except Exception as e:
        logger.error(f"Error en generar_reporte_word: {e}", exc_info=True)
        return f"Error al generar el archivo Word: {e}"

analytics_agent = adk.Agent(
    name="AnalyticsAgent",
    model="gemini-2.5-flash",
    instruction="Eres un analista de datos y BI. Tu propósito es generar resúmenes financieros, crear gráficos (<img>), exportar a Excel (<a>), consultar KPIs y generar reportes en formato PDF o Word. IMPORTANTE: Si el usuario te pide incluir un gráfico en un reporte PDF o Word, PRIMERO debes ejecutar la herramienta 'generar_grafico_analisis', leer el nombre del archivo generado en su respuesta (ej. grafico_xxx.png), y LUEGO ejecutar la herramienta de reporte pasándole ese nombre en el parámetro 'nombre_imagen'.",
    tools=[generar_grafico_analisis, obtener_resumen_pipeline, exportar_datos_excel, consultar_dashboard_bi, generar_reporte_pdf, generar_reporte_word]
)